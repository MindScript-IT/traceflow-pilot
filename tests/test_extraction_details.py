"""Regression tests for extraction fidelity bugs found in the pilot review."""
import io

from openpyxl import Workbook

from app.services.document_items import sanitize_doc_items, table_to_doc_items
from app.services.extraction import (
    coerce_pdf_location,
    csv_to_table,
    sanitize_items,
    xlsx_to_table,
)


def test_csv_semicolon_delimiter_is_detected():
    rows = csv_to_table(b"campo;valore\nsku;ABC-123\ncolore;nero\n")
    values = {(r["cell"], r["value"]) for r in rows}
    assert ("Row 2 / campo", "sku") in values
    assert ("Row 2 / valore", "ABC-123") in values
    assert ("Row 3 / campo", "colore") in values
    assert ("Row 3 / valore", "nero") in values


def test_csv_windows_1252_is_accepted():
    body = "campo;città\nsku;perché\n".encode("cp1252")
    rows = csv_to_table(body)
    assert any(r["cell"] == "Row 2 / città" for r in rows)
    assert any(r["value"] == "perché" for r in rows)


def test_xlsx_integral_floats_do_not_get_a_dot_zero():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 12345.0
    ws["A2"] = 0.25
    buf = io.BytesIO()
    wb.save(buf)
    values = [r["value"] for r in xlsx_to_table(buf.getvalue())]
    assert "12345" in values
    assert "12345.0" not in values
    assert "0.25" in values


def test_pdf_location_normalizes_page_abbreviations():
    assert coerce_pdf_location("Page 3") == "Page 3"
    assert coerce_pdf_location("p. 3") == "Page 3"
    assert coerce_pdf_location("pag. 4-5") == "Page 4-5"
    assert coerce_pdf_location("pagina 7") == "Page 7"
    assert coerce_pdf_location("Page three") is None


def test_passport_items_are_deduplicated():
    item = {"field_key": "sku", "value": "ABC", "unit": None, "location": "Page 1"}
    clean = sanitize_items([item, dict(item)], {"sku"})
    assert len(clean) == 1


def test_table_transcription_fallback_keeps_every_cell():
    table = [
        {"sheet": "Sheet 1", "cell": "A1", "value": "REACH"},
        {"sheet": "Sheet 1", "cell": "B1", "value": "pass"},
    ]
    items = table_to_doc_items(table)
    assert len(items) == 2
    assert items[0]["location"] == "Sheet 1 / A1"
    assert sanitize_doc_items(items)[0]["value"] == "REACH"
